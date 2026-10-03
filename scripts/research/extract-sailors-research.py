import hashlib
import json
from pathlib import Path
import openpyxl

source = Path('output/sailors-research.xlsx')
book = openpyxl.load_workbook(source, data_only=True)
candidates = []
verified_ids = {'Innocent': 59055, 'Realistic': 59060, 'Confident': 59061, 'Quick': 59057, 'Curious': 59067, 'Calculating': 59058, 'Experienced': 59066, 'Treasure-Seeking': 59059, 'Tenacious': 59062,
                'Ambitious': 59053, 'Diligent': 59054, 'Enamored': 59056, 'Honest': 59063, 'Tough': 59064, 'Strong': 59065,
                'Dreaming of a Full Haul': 59068, 'Born on the Sea': 59070, 'Powerful': 59069, 'Smart': 59071,
                'Quick-Witted': 59072, 'Arkahn': 59101, 'Hetario': 59227, 'Pacuna': 59228}
for sheet in book:
    if sheet.title in ['Info', 'All Sailors', 'My Sailors V1.0'] or sheet.title.startswith('New '):
        continue
    candidates.append({'sourceName': sheet.title, 'gameId': verified_ids.get(sheet.title),
                       'regionValidation': 'pending-SA', 'sourceSheet': sheet.title})
puro = book['Innocent']
ranges = []
for index in range(8):
    lower = puro.cell(22 + index, 13).value
    upper = puro.cell(32 + index, 13).value
    assert isinstance(lower, (float, int)) and isinstance(upper, (float, int))
    assert 0 <= lower <= upper
    min_growth = sum(puro.cell(2 + index, col).value for col in range(4, 14))
    max_growth = sum(puro.cell(12 + index, col).value for col in range(4, 14))
    assert abs(min_growth - lower) < 0.000001 and abs(max_growth - upper) < 0.000001
    ranges.append({'sourceAttribute': puro.cell(22 + index, 3).value, 'minimum': lower, 'maximum': upper,
                   'minimumCell': f'M{22 + index}', 'maximumCell': f'M{32 + index}'})
record = {'retrievedAt': '2026-10-03', 'level': 10, 'status': 'community-reference-not-SA-validated',
          'sourceUrl': 'https://docs.google.com/spreadsheets/d/1CFJOgyhnw2_Rq4UM2zs2uK6J15nipv1LiwrDQUWxU0o/edit',
          'sourceTitle': 'BDO Sailors 2024 - Public v1.0', 'sha256': hashlib.sha256(source.read_bytes()).hexdigest(),
          'candidates': candidates, 'puro': {'gameId': 59055, 'sourceSheet': 'Innocent', 'ranges': ranges}}
Path('docs/marinheiros-rag').mkdir(exist_ok=True)
Path('docs/marinheiros-rag/registro.json').write_text(json.dumps(record, ensure_ascii=False, indent=2), encoding='utf-8')
print(f'{len(candidates)} candidatos registrados; 8 limites do Puro conferidos contra somas de crescimento.')
